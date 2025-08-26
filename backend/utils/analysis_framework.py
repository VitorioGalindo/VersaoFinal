from backend.models import Company
from docx import Document
import os

def generate_analysis_framework(company: Company) -> str:
    """
    Gera um framework de análise padrão para uma empresa a partir de um template .docx.
    
    Args:
        company (Company): Objeto da empresa com os dados.
        
    Returns:
        str: Texto com o framework de análise preenchido.
    """
    try:
        # Caminho para o arquivo de template
        template_path = os.path.join('frameworks', 'acoes-ronan.docx')
        
        # Verificar se o arquivo de template existe
        if not os.path.exists(template_path):
            # Se não existir, usar um template padrão
            framework = f"""# Análise de {company.company_name} ({company.b3_issuer_code if company.b3_issuer_code else 'N/A'})

## Resumo Executivo
- Setor: {company.b3_sector if company.b3_sector else 'N/A'}
- Subsetor: {company.b3_subsector if company.b3_subsector else 'N/A'}
- Segmento: {company.b3_segment if company.b3_segment else 'N/A'}
- Data de Fundação: {company.founded_date.strftime('%d/%m/%Y') if company.founded_date else 'N/A'}
- Número de Funcionários: {company.employee_count if company.employee_count else 'N/A'}

## Atividade Principal
{company.main_activity if company.main_activity else 'N/A'}

## Visão Geral Financeira
(Adicionar análise dos principais indicadores financeiros)

## Análise de Documentos
(Adicionar análise dos documentos registrados na CVM)

## Riscos e Oportunidades
(Adicionar análise dos principais riscos e oportunidades)

## Recomendação
(Adicionar recomendação de investimento)

## Referências
- Documentos CVM: [Link para documentos]
- Website: {company.website if company.website else 'N/A'}
"""
            return framework
        
        # Ler o documento .docx
        doc = Document(template_path)
        
        # Converter o conteúdo do documento para texto, preservando a estrutura
        full_text = []
        for paragraph in doc.paragraphs:
            # Obter o texto do parágrafo
            text = paragraph.text
            
            # Verificar se o parágrafo tem numeração
            if paragraph._element.xpath('.//w:numPr'):
                # Se tiver numeração, manter o texto como está
                full_text.append(text)
            else:
                # Se não tiver numeração, manter o texto como está
                full_text.append(text)
        
        framework = '\n'.join(full_text)
        
        # Substituir placeholders no template (se houver)
        framework = framework.replace('{company_name}', company.company_name or 'N/A')
        framework = framework.replace('{ticker}', company.b3_issuer_code or 'N/A')
        framework = framework.replace('{sector}', company.b3_sector or 'N/A')
        framework = framework.replace('{subsector}', company.b3_subsector or 'N/A')
        framework = framework.replace('{segment}', company.b3_segment or 'N/A')
        framework = framework.replace(
            '{founded_date}', 
            company.founded_date.strftime('%d/%m/%Y') if company.founded_date else 'N/A'
        )
        framework = framework.replace(
            '{employees}', 
            str(company.employee_count) if company.employee_count else 'N/A'
        )
        framework = framework.replace('{activity}', company.main_activity or 'N/A')
        framework = framework.replace('{website}', company.website or 'N/A')
        
        return framework
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        # Em caso de erro, usar o template padrão
        framework = f"""# Análise de {company.company_name} ({company.b3_issuer_code if company.b3_issuer_code else 'N/A'})

## Resumo Executivo
- Setor: {company.b3_sector if company.b3_sector else 'N/A'}
- Subsetor: {company.b3_subsector if company.b3_subsector else 'N/A'}
- Segmento: {company.b3_segment if company.b3_segment else 'N/A'}
- Data de Fundação: {company.founded_date.strftime('%d/%m/%Y') if company.founded_date else 'N/A'}
- Número de Funcionários: {company.employee_count if company.employee_count else 'N/A'}

## Atividade Principal
{company.main_activity if company.main_activity else 'N/A'}

## Visão Geral Financeira
(Adicionar análise dos principais indicadores financeiros)

## Análise de Documentos
(Adicionar análise dos documentos registrados na CVM)

## Riscos e Oportunidades
(Adicionar análise dos principais riscos e oportunidades)

## Recomendação
(Adicionar recomendação de investimento)

## Referências
- Documentos CVM: [Link para documentos]
- Website: {company.website if company.website else 'N/A'}
"""
        return framework